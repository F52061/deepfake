from pathlib import Path

from docx import Document
from lxml import etree


workspace = Path(__file__).resolve().parents[2]
document = Document(workspace / "跨域鉴伪_项目结题报告书_初稿.docx")
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.startswith(("B3.", "C.", "D.", "E.")):
        print(index, paragraph.text, "BREAKS", [node.attrib for node in paragraph._p.xpath(".//w:br")])
for section in document.sections:
    print("SECTION", section.start_type)
for section in document._element.xpath(".//w:sectPr"):
    print("SECTION PARENT", section.getparent().tag, section.getparent().getparent().xpath(".//w:t/text()")[:3])
for index, table in enumerate(document.tables):
    print("TABLE", index, table.cell(0, 0).text, etree.tostring(table._tbl.tblPr, encoding="unicode").split("</w:tblPr>")[0][-900:])

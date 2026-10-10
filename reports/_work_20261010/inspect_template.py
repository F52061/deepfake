import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

workspace = Path(__file__).resolve().parents[2]
reference = workspace / "红字批注版附件2.参赛作品信息暨立项项目结题报告书.docx"
document = Document(reference)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text:
        print("P", index, repr(paragraph.text), paragraph.style.name)
        print("pPr", etree.tostring(paragraph._p.pPr, encoding="unicode").split(">")[1:] if paragraph._p.pPr is not None else None)
        print("fonts", [(run.font.name, run.font.size.pt if run.font.size else None, run.bold) for run in paragraph.runs[:2]])
for table_index, table in enumerate(document.tables):
    print("TABLE", table_index, "rows", len(table.rows), "columns", len(table.columns))
    print("grid", [column.width for column in table.columns])
    for row_index, row in enumerate(table.rows):
        print(row_index, [(cell_index, cell.text) for cell_index, cell in enumerate(row.cells)])
inventory = {}
with ZipFile(reference) as archive:
    for name in archive.namelist():
        content = archive.read(name)
        inventory[name] = {"bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()}
    body = etree.fromstring(archive.read("word/document.xml"))
    print("SECTIONS", [etree.tostring(section, encoding="unicode") for section in body.xpath("//w:sectPr", namespaces=ns)])
    print("FIELDS", body.xpath("//w:instrText/text()", namespaces=ns))
    print("CONTROLS", len(body.xpath("//w:sdt", namespaces=ns)))
Path(__file__).with_name("template_inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
Path(__file__).with_name("reference.sha256").write_text(hashlib.sha256(reference.read_bytes()).hexdigest(), encoding="utf-8")

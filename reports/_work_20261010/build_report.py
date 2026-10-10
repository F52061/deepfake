import hashlib
import io
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


WORKSPACE = Path(__file__).resolve().parents[2]
TASK_DIR = Path(__file__).resolve().parent
REFERENCE = WORKSPACE / "红字批注版附件2.参赛作品信息暨立项项目结题报告书.docx"
OUTPUT = WORKSPACE / "跨域鉴伪_项目结题报告书_初稿.docx"
PROJECT = "跨域鉴伪——面向社媒的人脸深度伪造检测系统"
MEMBERS = "黄耀阳、张晶、魏家冰、刘华威、周智浩、胡懿珊、梁建宁"
BLACK = RGBColor(0, 0, 0)
NAMESPACES = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def format_run(run, size=12, font="宋体", bold=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.font.bold = bold
    run.font.underline = False
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), font)


def format_paragraph(paragraph, indent=False, spacing=1.5, after=4):
    paragraph.paragraph_format.line_spacing = spacing
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(24 if indent else 0)
    paragraph.paragraph_format.widow_control = True
    snap = paragraph._p.get_or_add_pPr().find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        paragraph._p.get_or_add_pPr().append(snap)
    snap.set(qn("w:val"), "0")


def fill_cell(cell, text, size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, content in enumerate(text.split("\n")):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = align
        format_paragraph(paragraph, spacing=1.25, after=3)
        format_run(paragraph.add_run(content), size=size)


def table_format(table, keep_rows=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + side))
        if element is None:
            element = OxmlElement("w:" + side)
            borders.append(element)
        for key, value in {"val": "single", "sz": "6", "color": "D9D9D9"}.items():
            element.set(qn("w:" + key), value)
    for row in table.rows:
        row.height = None
        row_properties = row._tr.get_or_add_trPr()
        for height in list(row_properties.findall(qn("w:trHeight"))):
            row_properties.remove(height)
        if keep_rows:
            element = OxmlElement("w:cantSplit")
            row_properties.append(element)
        for cell in row.cells:
            margin = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            if margin is None:
                margin = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(margin)
            for side in ("top", "bottom", "left", "right"):
                element = margin.find(qn("w:" + side))
                if element is None:
                    element = OxmlElement("w:" + side)
                    margin.append(element)
                element.set(qn("w:w"), "90")
                element.set(qn("w:type"), "dxa")


def body_text(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    format_paragraph(paragraph, indent=True)
    format_run(paragraph.add_run(text))
    return paragraph


def heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    format_paragraph(paragraph, after=7)
    paragraph.paragraph_format.space_before = Pt(10)
    format_run(paragraph.add_run(text), size=14 if level == 1 else 12, font="黑体", bold=True)
    return paragraph


def caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(paragraph, after=8)
    format_run(paragraph.add_run(text), size=10.5)


def data_table(document, labels, rows, widths):
    table = document.add_table(rows=1, cols=len(labels))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, label in enumerate(labels):
        fill_cell(table.rows[0].cells[index], label, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EDEDED")
        table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shading)
        for run in table.rows[0].cells[index].paragraphs[0].runs:
            run.bold = True
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            fill_cell(cells[index], str(value), size=10.5,
                      align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    table_format(table, keep_rows=True)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def font(size, bold=False):
    path = Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc")
    return ImageFont.truetype(str(path), size)


def centered_text(draw, position, text, size=28, bold=False):
    draw.multiline_text(position, text, font=font(size, bold), fill="black", anchor="mm", align="center", spacing=8)


def make_figures(acc_rows):
    diagram = Image.new("RGB", (1600, 760), "white")
    draw = ImageDraw.Draw(diagram)
    draw.text((35, 20), "当前工程原型", font=font(36, True), fill="black")
    boxes = [(40, 95, 335, 205, "前端界面\n选择图片"),
             (435, 95, 730, 205, "后端推理\n改进单 ViT"),
             (830, 95, 1125, 205, "真假分类\n返回检测结果"),
             (1225, 95, 1555, 205, "界面展示\n供人工参考")]
    for left, top, right, bottom, label in boxes:
        draw.rectangle((left, top, right, bottom), outline="#555555", width=3)
        centered_text(draw, ((left + right) / 2, (top + bottom) / 2), label, 30)
    for start in (335, 730, 1125):
        draw.line((start + 8, 150, start + 90, 150), fill="black", width=3)
        draw.polygon([(start + 90, 150), (start + 76, 142), (start + 76, 158)], fill="black")
    draw.line((35, 255, 1560, 255), fill="#bbbbbb", width=2)
    draw.text((35, 290), "独立研究验证路线", font=font(36, True), fill="black")
    research_boxes = [(40, 425, 290, 555, "同一图像"),
                      (410, 360, 740, 475, "改进 ViT\nCLS 与多层 tokens"),
                      (410, 550, 740, 665, "CLIP image\nCLS 与多层 tokens"),
                      (860, 420, 1190, 605, "全局读取与\nBridge 多层融合\n及受控对照"),
                      (1290, 420, 1560, 605, "分类评测\n收益与误伤分析")]
    for left, top, right, bottom, label in research_boxes:
        draw.rectangle((left, top, right, bottom), outline="#555555", width=3)
        centered_text(draw, ((left + right) / 2, (top + bottom) / 2), label, 28)
    for coordinates in [((290, 490), (345, 490), (345, 417), (410, 417)),
                        ((345, 490), (345, 607), (410, 607)),
                        ((740, 417), (800, 417), (800, 470), (860, 470)),
                        ((740, 607), (800, 607), (800, 555), (860, 555)),
                        ((1190, 512), (1290, 512))]:
        draw.line(coordinates, fill="black", width=3)
        tip = coordinates[-1]
        draw.polygon([tip, (tip[0] - 14, tip[1] - 8), (tip[0] - 14, tip[1] + 8)], fill="black")
    draw.text((40, 715), "注  研究用融合路线未作为当前界面后端的已验证部署方案", font=font(26), fill="black")
    diagram.save(TASK_DIR / "architecture.png")
    chart = Image.new("RGB", (1500, 790), "white")
    draw = ImageDraw.Draw(chart)
    draw.text((65, 20), "改进单 ViT 阶段训练记录", font=font(36, True), fill="black")
    left, top, right, bottom = 125, 110, 1430, 645
    for tick in range(75, 101, 5):
        vertical = bottom - (tick - 75) / 25 * (bottom - top)
        draw.line((left, vertical, right, vertical), fill="#dddddd", width=2)
        draw.text((55, vertical - 14), str(tick), font=font(27), fill="black")
    draw.line((left, top, left, bottom, right, bottom), fill="black", width=3)
    for epoch in (15, 20, 25, 30, 35, 40, 45, 50):
        horizontal = left + (epoch - 15) / 35 * (right - left)
        draw.line((horizontal, bottom, horizontal, bottom + 8), fill="black", width=2)
        draw.text((horizontal - 19, bottom + 15), str(epoch), font=font(26), fill="black")
    for key, color, label in [("ffpp", "#245a84", "FF++ AUC"), ("cross3", "#a14b29", "CD2 / DFDCP / Wild 等权平均 AUC")]:
        positions = [(left + (row["epoch"] - 15) / 35 * (right - left),
                      bottom - (row[key] - 75) / 25 * (bottom - top)) for row in acc_rows]
        draw.line(positions, fill=color, width=5)
        draw.ellipse((positions[-1][0] - 6, positions[-1][1] - 6,
                      positions[-1][0] + 6, positions[-1][1] + 6), fill=color)
        legend_top = 65
        legend_left = 155 if key == "ffpp" else 525
        draw.line((legend_left, legend_top + 15, legend_left + 75, legend_top + 15), fill=color, width=5)
        draw.text((legend_left + 90, legend_top), label, font=font(27), fill="black")
    draw.text((65, 735), "AUC 单位为百分比；完整展示 epoch 15 至 50，不据目标域曲线挑选最优模型。", font=font(25), fill="black")
    chart.save(TASK_DIR / "training_auc.png")


def main():
    if OUTPUT.exists():
        prior_check = json.loads((TASK_DIR / "build_checks.json").read_text(encoding="utf-8"))
        if hashlib.sha256(OUTPUT.read_bytes()).hexdigest() != prior_check["output_sha256"]:
            raise FileExistsError(f"Output has changed since this task created it: {OUTPUT}")
    expected_hash = (TASK_DIR / "reference.sha256").read_text(encoding="utf-8")
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == expected_hash
    acc_rows = []
    for line in (WORKSPACE / "acc.txt").read_text(encoding="utf-8").splitlines():
        epoch_match = re.search(r"EPOCH=(\d+)", line)
        metrics = dict(re.findall(r"(\w+)\|acc:[\d.]+%, auc:([\d.]+)%", line))
        if epoch_match and metrics:
            values = {key: float(value) for key, value in metrics.items()}
            values["epoch"] = int(epoch_match.group(1))
            values["cross3"] = (values["cd2"] + values["dfdcp"] + values["wilddf"]) / 3
            acc_rows.append(values)
    assert len(acc_rows) == 36 and acc_rows[-1]["epoch"] == 50
    make_figures(acc_rows)
    source_summary = json.loads((WORKSPACE / "vit_module/_g28/clip_readout_run01/summary.json").read_text(encoding="utf-8"))
    macro = source_summary["primary_macro"]
    pooled = macro["SELECTED|POOLED|ensemble_vs_V_BASE"]
    assert abs(pooled["baseline_macro_auc"] - 0.8328148148148148) < 1e-12
    document = Document(REFERENCE)
    tables = list(document.tables)
    paragraphs = list(document.paragraphs)
    original_blocks = list(document._element.body)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name] if style_name in document.styles else document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.color.rgb = BLACK
        style.font.name = "黑体"
        style.font.underline = False
    for paragraph in paragraphs:
        if paragraph.text.startswith("B1-3") or paragraph.text.startswith("（成果分类") or paragraph.text.startswith("（各"):
            paragraph._p.getparent().remove(paragraph._p)
    paragraphs[0].style = document.styles["Title"]
    for run in paragraphs[0].runs:
        run.font.color.rgb = BLACK
    paragraphs[6].runs[0].text = "2026年10月"
    cover = tables[0]
    cover_values = [PROJECT, "【待补充项目编号】", "科技发明制作类（A／B具体类别待确认）", "信息技术",
                    "【待补充立项类别】", "黄耀阳", "计算机科学与技术2024级", "15017260728",
                    "张健、周志立、陈淑红", "计算机科学与网络工程学院"]
    for row, value in zip(cover.rows, cover_values):
        fill_cell(row.cells[1], value)
    table_format(cover, keep_rows=True)
    author = tables[1]
    for row_index, cell_index, value in [(0, 2, "黄耀阳"), (0, 7, "男"), (0, 11, "2006年3月"),
                                         (1, 2, "计算机科学与技术2024级【待补充班级】"),
                                         (2, 2, "本科"), (2, 7, "4年"), (2, 11, "2024年9月"),
                                         (3, 4, PROJECT), (4, 4, "□个人项目  ☑集体项目"), (5, 4, "无")]:
        fill_cell(author.rows[row_index].cells[cell_index], value, align=WD_ALIGN_PARAGRAPH.CENTER)
    assignments = [("黄耀阳", "改进ViT与注意力抑制实现"), ("张晶", "特征融合与实验验证"),
                   ("魏家冰", "模型架构搭建与训练"), ("刘华威", "后端接口与模型接入"),
                   ("周智浩", "系统优化与联调"), ("胡懿珊", "技术调研与文书整理"),
                   ("梁建宁", "前端页面设计")]
    for index, (name, assignment) in enumerate(assignments, start=7):
        row = author.rows[index]
        for cell_index, value in [(1, "【待补学号】"), (3, name),
                                 (6, "计算机科学与网络工程学院" if index == 7 else "【待确认学院】"),
                                 (8, "本科"), (10, assignment)]:
            fill_cell(row.cells[cell_index], value, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in list(author.rows)[14:]:
        author._tbl.remove(row._tr)
    table_format(author, keep_rows=True)
    for row in author.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph, spacing=1.15, after=1)
                for run in paragraph.runs:
                    format_run(run, size=11)
    b3 = tables[2]
    basic = (
        "目的与思路：面向社媒人脸图像鉴伪，研制可运行的图片真假检测原型，并评估跨数据集泛化能力。实际研究以改进ViT为检测主体，探索CLIP视觉表征及Bridge多层融合，未将原申报的空域—频域双流方案全部实现为最终系统。\n"
        "技术特点：迁移应用高响应注意力随机抑制，探索补丁分组与特征一致性训练；将改进ViT接入参考M2F2_Det的融合框架；利用冻结特征、噪声重训及其他视频替换对照检查CLIP信息增量。上述设计与已有方法的关系及性能限制见E部分。\n"
        "主要指标：单ViT第50轮记录为FF++ AUC 98.162%，CD2 86.468%，DFDCP 80.893%，Wild 79.663%，FFIW 72.470%。这是同一轮训练记录，不是逐域择优结果，亦不等于线上检测成功率。前端与后端已接入改进单ViT，可完成图片真假鉴别；高逼真伪造检测仍不及预期。"
    )
    science = (
        "作品采用Transformer特征提取与训练正则化机制，围绕局部响应依赖和融合信息利用进行实验验证。相较仅报告总体成绩，本项目进一步区分原生ViT分类头、ViT-only探针和完整融合检测器，通过同输入对照及视频级统计定位融合收益与损伤。\n"
        "已定位原Bridge压缩层LayerNorm(1)造成输入信息消失的实现缺陷，但修复尝试尚未证明稳定性能提升。冻结特征复验中，全局CLIP读取跨域宏平均点估计为84.074%，ViT-only探针为83.281%；增益区间包含零，不能认定显著超越。\n"
        "项目特色在于改进ViT的系统集成与融合瓶颈的受控诊断，尚无证据支持全面领先现有方法或完成泛化难题的解决。参考文献及取数说明见E部分。"
    )
    usage = (
        "使用流程：通过前端选择图片，后端调用改进单ViT模型，返回真假判断并在界面展示。当前部署原型不使用研究用CLIP／Bridge融合检测器。\n"
        "适用范围：教学科研、实验样本演示和辅助筛查。检测输出不能单独作为金融核验、司法取证或自动内容处置的依据；对未见伪造算法、高逼真合成及传播降质样本的可靠性仍需专项验证。\n"
        "推广前景：可在进一步验证后作为人工复核的辅助组件。项目尚无商业客户、实际业务收益或规模化部署证据，不申报并发吞吐、毫秒级时延或经济效益达标。"
    )
    b3_values = [PROJECT, "B．信息技术", basic, science, "无已取得的获奖或鉴定成果申报。", "A．实验室阶段",
                 "暂无技术转让；后续合作应履行学校及相关权利人的审批程序。",
                 "☑模型  ☑现场演示  ☑图片\n展示内容：检测原型、架构示意图及实验结果。", usage, "☑未提出专利申请"]
    for row, value in zip(b3.rows, b3_values):
        fill_cell(row.cells[1], value, size=11.5)
        for paragraph in row.cells[1].paragraphs:
            format_paragraph(paragraph, spacing=1.2, after=2)
        for paragraph in row.cells[0].paragraphs:
            format_paragraph(paragraph, spacing=1.2, after=2)
            for run in paragraph.runs:
                format_run(run, size=11.5)
    table_format(b3)
    current_research = (
        "深度伪造检测已形成基于卷积纹理、Transformer结构建模、频域统计及预训练视觉表征的多种技术路线。FaceForensics++和Celeb-DF等基准用于评估训练域内能力与跨数据集迁移，但受控数据集成绩不能直接代表真实社媒场景的可用性[1][2]。\n"
        "ViT通过补丁序列建模图像结构，CLIP利用图文对比预训练形成视觉表征，两者训练目标不同，但可能提取相关的图像信息[3][4]。将两路特征拼接或增加融合层，并不必然产生可迁移增益，需要以独立ViT、简单读取和容量匹配对照加以验证。\n"
        "本项目参考既有M2F2_Det工程的多层融合思路，将原检测分支替换为改进ViT，并检查训练增强与融合结构的实际作用。阶段实验发现，全局读取具有内容依赖线索，局部自适应读取尚未证明稳定超过基线；这说明后续研究需兼顾特征互补性、源域选参迁移及对原有ViT判别能力的保护，而不能仅凭结构复杂度判断先进性。\n"
        "现阶段作品属于模型实现与实验验证原型。项目没有完成与所有同类方法在统一协议下的性能排名，也没有确认重度压缩鲁棒性、轻量化或高并发指标达标。"
    )
    fill_cell(tables[3].cell(0, 0), current_research)
    table_format(tables[3])
    d_heading = next(paragraph for paragraph in paragraphs if paragraph.text == "D.项目主要成果")
    e_heading = next(paragraph for paragraph in paragraphs if paragraph.text == "E.参赛作品")
    body = document._element.body
    d_position = list(body).index(d_heading._p)
    e_position = list(body).index(e_heading._p)
    d_source_heading = deepcopy(paragraphs[24]._p)
    other_template = deepcopy(tables[8]._tbl)
    for element in list(body)[d_position + 1:e_position]:
        body.remove(element)
    insertion_point = d_heading._p
    for content in ["一、论文发表情况：无。", "二、专利情况：未提出本项目专利申请，无授权专利。",
                    "三、获奖及采访情况：本期不申报相关成果。", "四、其他已完成成果"]:
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, after=5)
        format_run(paragraph.add_run(content), font="黑体", bold=True)
        insertion_point.addnext(paragraph._p)
        insertion_point = paragraph._p
    from docx.table import Table
    for name, evidence in [
        ("改进ViT模型及图片真假检测前后端原型", "模型实现、单ViT训练记录；前端调用后端并展示真假结果。属于实验室原型，尚未形成业务级验证。"),
        ("跨域评测诊断工具与阶段研究资料", "形成特征读取、残差及局部读取、内容替换和视频级统计脚本，以及训练记录、融合缺陷分析、控制实验汇总与研究报告。未作为已发表论文申报。")]:
        copied_table = deepcopy(other_template)
        insertion_point.addnext(copied_table)
        inserted_table = Table(copied_table, document._body)
        for row, value in zip(inserted_table.rows, [name, PROJECT, "项目阶段性成果，不申报国家级或省部级认定", "项目组自主研发，无外部成果认定机构", MEMBERS]):
            fill_cell(row.cells[1], value, size=11)
        extra_row = inserted_table.add_row()
        fill_cell(extra_row.cells[0], "成果内容", size=11)
        fill_cell(extra_row.cells[1], evidence, size=11)
        for row in inserted_table.rows:
            for paragraph in row.cells[0].paragraphs:
                format_paragraph(paragraph, spacing=1.25)
                for run in paragraph.runs:
                    format_run(run, size=11)
        table_format(inserted_table, keep_rows=True)
        insertion_point = copied_table
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)
        insertion_point.addnext(spacer._p)
        insertion_point = spacer._p
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(("A.", "B3.", "C.", "D.", "E.")):
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            for page_break in paragraph._p.xpath('.//w:br[@w:type="page"]'):
                page_break.getparent().remove(page_break)
    for paragraph in paragraphs[20:23]:
        if paragraph._p.getparent() is not None and not paragraph.text:
            paragraph._p.getparent().remove(paragraph._p)
    seen_b3 = False
    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("B3."):
            seen_b3 = True
        if seen_b3 and not paragraph.text.strip() and paragraph._p.xpath('.//w:br[@w:type="page"]') and not paragraph._p.xpath('.//w:sectPr'):
            paragraph._p.getparent().remove(paragraph._p)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(title, after=8)
    format_run(title.add_run(PROJECT), size=16, font="黑体", bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(subtitle, after=8)
    format_run(subtitle.add_run("项目研究报告"), size=14, font="黑体", bold=True)
    caption(document, "研究与成果统计截至2026年10月10日")
    heading(document, "摘要")
    body_text(document, "本项目针对社媒人脸图像的深度伪造鉴别需求，完成了改进ViT模型的训练评测，构建了可调用该模型的前端与后端原型，并探索CLIP视觉编码器及Bridge多层融合。单ViT阶段记录显示，模型在FF++具有较好的区分能力，但跨数据集和高逼真伪造样本的检测仍不及预期。针对融合未产生稳定增益的问题，我们开展了信息传递检查、全局与局部特征读取、噪声重训和其他视频替换实验，定位了实现缺陷并限定了现阶段可支持的研究结论。项目已形成检测原型、验证工具和分析材料，但尚未达到申报书提出的完整性能与成果目标。")
    heading(document, "一 研究目标与实际实施")
    body_text(document, "社媒传播中的人脸图片来源复杂，伪造方法、拍摄条件和后处理流程与训练集可能存在差异。只在已知伪造类型上取得较高检测成绩，不能保证模型面对新的图片仍能可靠工作。本项目希望减少检测器对局部显性痕迹的依赖，并检验不同预训练表征能否补充可迁移的真假判别信息。")
    body_text(document, "原申报书规划空域—频域双流、注意力抑制与补丁重组、跨域和降质验证，以及标准化服务建设。实际实施阶段，我们将主要研究资源用于改进ViT及其与CLIP视觉表征的融合探索。CLIP不是频域编码器，因此当前融合框架不等同于原计划的DCT频域双流方案。前后端实现了基础图片鉴别功能，但双域协同去偏、完整压缩鲁棒性目标、部署性能指标和论文专利成果没有全部完成。")
    body_text(document, "研究过程中，我们从以总体AUC为主的评测逐步转向受控诊断，分别检查模型分支、分类读取方式和训练选参是否造成跨域损伤。这样的调整使项目获得了关于当前方法边界的具体证据，但不能据此声称已解决跨域泛化问题，或已经证明模型存在不可突破的理论上限。")
    data_table(document, ["原计划事项", "实际进展", "结题口径"], [
        ["检测模型与训练增强", "改进ViT实现及训练评测", "已形成阶段实现"],
        ["双流特征融合", "转向ViT／CLIP及Bridge验证", "路线调整，增益未确认"],
        ["跨域与鲁棒性验证", "多域评测及受控扰动探索", "未达到全部预期指标"],
        ["服务与界面", "前端、后端接入改进单ViT", "基础原型已实现"],
        ["论文与专利", "无本项目论文专利成果", "未完成"]], [1.45, 2.2, 1.9])
    heading(document, "二 检测模型与研究框架")
    heading(document, "改进ViT检测分支", level=2)
    body_text(document, "独立检测器采用12层ViT，将人脸图片表示为补丁序列，以分类token的高维特征完成真假判断。训练设计引入高响应注意力随机抑制：在部分训练前向中选取较高响应位置，以其余位置的平均响应替换部分选中项，再归一化注意力。目的是减少少数局部响应主导训练的倾向。该模块是训练期机制，常规评测时关闭随机抑制，不应将其描述为推理期反复随机鉴别。")
    body_text(document, "模型代码还提供依据分类token注意力进行补丁分组、交叉组合并提取部分序列表征的路径，用于探索部分区域与完整图像表征的一致性。上述实现沿用并迁移已有注意力抑制与补丁重组思想。当前单ViT成绩不能单独证明每个训练机制均有独立增益，还需要配平的消融实验。")
    heading(document, "CLIP与Bridge研究用融合框架", level=2)
    body_text(document, "研究框架参考M2F2_Det的多层融合结构，将检测分支替换为改进ViT，保留CLIP image编码器。两路首先独立编码图像：ViT输出768维CLS，CLIP取倒数第二层的1024维CLS；不同维度经投影后进入分类空间。Bridge另行读取两路的中间层patch tokens，降到64维后沿序列维拼接，并将前一级融合序列加入后一级，经过三级Transformer形成128维紧凑表征。最终分类头读取ViT、CLIP CLS和Bridge输出。")
    body_text(document, "当前实现的ViT钩子为blocks[3]、blocks[6]、blocks[9]，按一基计数对应第4、7、10层；CLIP取第4、8、23层附近的输出。这里说明的是实际代码，不把两路统称为第3、6、9层。Bridge在下游融合，不把CLIP信息回写到ViT CLS。Phase 1融合训练加载既有ViT权重并冻结两路编码器，重点更新投影、融合和分类模块。")
    body_text(document, "当前前后端实际调用的是改进单ViT，而不是该研究融合检测器。CLIP text在原Bridge实现中虽被计算，但没有接入最终真假分类的有效路径；本期有关CLIP增量的实验主要检验视觉编码器，不能将其称为已经实现图文语义联合鉴伪。")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(TASK_DIR / "architecture.png"), width=Inches(5.45))
    caption(document, "图1 当前单ViT工程原型与研究用融合路线示意")
    heading(document, "三 数据与评测方法")
    body_text(document, "模型围绕FF++开展训练，在CD2、DFDCP、Wild及FFIW等数据集评估。AUC反映真假样本分数排序的区分能力，准确率依赖固定分类规则，两者不能互换。FF++域内成绩、跨域逐集成绩以及跨域宏平均分别报告，避免用一个聚合数字掩盖不同目标集的差异。")
    body_text(document, "本报告区分两个取数口径。第一类为单ViT训练日志，展示同一epoch在各评测集的记录；第二类为G28冻结特征读取实验，采用源域训练的ViT-only线性探针为基线，其分类头不同于单ViT原生头。G28主要宏平均仅包含CD2、DFDCP、Wild；CD1与CD2存在包含关系，单列而不重复加权，FFIW不计入该主指标。两类表格不能直接相减来估计模块贡献。")
    body_text(document, "G28采用源域视频分组交叉验证选取读取器参数，使用视频级配对bootstrap评估差异，三个随机种子与内容替换重复分别记录。诊断数据将1标为伪造；历史模型输出方向另行核验，不能仅凭logit下标判断真假。目标域已被多轮查看，且历史融合检查点曾依据FFIW结果选模，因此本报告将相关发现作为探索性证据，不声称完全独立的确认试验。")
    heading(document, "四 阶段检测结果")
    heading(document, "独立改进ViT训练记录", level=2)
    body_text(document, "表1取现有acc.txt的第50轮完整记录，未按每个目标集分别挑选最佳epoch。这些读数说明模型在受控基准上能够进行鉴别，同时显示跨域性能明显低于FF++。训练日志没有提供本表全部评测集的样本数、逐样本预测及其与当前部署权重的对应证明，因此本表用于阶段实验展示，不将其解释为在线系统的已验收性能。")
    last_row = acc_rows[-1]
    accuracies = dict(re.findall(r"(\w+)\|acc:([\d.]+)%", (WORKSPACE / "acc.txt").read_text(encoding="utf-8").splitlines()[-1]))
    data_table(document, ["数据集", "准确率 %", "AUC %"],
               [[label, accuracies[key], f"{last_row[key]:.3f}"] for label, key in
                [("FF++", "ffpp"), ("CD2", "cd2"), ("DFDCP", "dfdcp"), ("Wild", "wilddf"), ("FFIW", "ffiw")]], [2.2, 1.65, 1.7])
    caption(document, "表1 独立改进ViT第50轮记录")
    body_text(document, f"CD2、DFDCP、Wild三域AUC的等权平均为{last_row['cross3']:.3f}%。这一平均不含FF++和FFIW，不能与其他包含不同数据集的历史平均混用。图2完整展示第15至50轮FF++及同口径三域平均AUC：后期域内成绩维持较高水平，但跨域没有同步持续改善。该现象说明仅延长训练不能保证迁移收益，并不构成性能上限的证明。")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(TASK_DIR / "training_auc.png"), width=Inches(5.45))
    caption(document, "图2 单ViT训练日志曲线  数据来源为acc.txt")
    heading(document, "冻结特征融合复验", level=2)
    comparisons = [("ViT-only线性探针", None), ("原融合检测器参考", "DETECTOR_REFERENCE_vs_V_BASE"),
                   ("ViT／CLIP线性拼接", "V_C_LINEAR_vs_V_BASE"),
                   ("全局CLIP修正 POOLED", "SELECTED|POOLED|ensemble_vs_V_BASE"),
                   ("局部自适应 ADAPTIVE", "SELECTED|ADAPTIVE|ensemble_vs_V_BASE")]
    result_rows = []
    for label, key in comparisons:
        if key is None:
            result_rows.append([label, f"{pooled['baseline_macro_auc'] * 100:.3f}", "基线"])
        else:
            entry = macro[key]
            result_rows.append([label, f"{entry['macro_auc'] * 100:.3f}",
                                f"{entry['delta_macro_auc'] * 100:+.3f}\n[{entry['ci95'][0] * 100:+.3f}, {entry['ci95'][1] * 100:+.3f}]"])
    data_table(document, ["读取方式", "主宏平均 AUC %", "差值及95%区间\n单位为百分点"], result_rows, [2.25, 1.5, 1.8])
    caption(document, "表2 G28探索性复验  三域宏平均  多种子模型为分数平均集成")
    body_text(document, "G28基线为83.281%，全局CLIP修正点估计为84.074%，差值为+0.793个百分点，95%区间为[−0.284，+1.994]个百分点，包含零。线性拼接和局部自适应读取的差异区间也包含零。因此，当前源域选参流程没有证明融合方案稳定优于纯ViT探针，更不能把表2当作相对部署单ViT的提升。不同读取器的点估计大小只提供下一步验证线索。")
    heading(document, "五 融合瓶颈及控制实验")
    heading(document, "信息传递实现缺陷", level=2)
    body_text(document, "原Bridge压缩模块将每个token从64维映射成一个标量，再使用LayerNorm(1)。单元素归一化的均值等于自身、方差为零，输出只剩可学习偏置，与图像内容无关。历史检查中，不同图片的Bridge嵌入逐元素相同，相关前级参数没有有效梯度。这解释了原模型虽包含多层融合结构，最终却主要依靠两路CLS分类的现象。")
    body_text(document, "该结果只针对当前实现，不否定多层融合的一般价值。修复与重新训练对照没有取得明确的稳定增益，说明消除工程缺陷只是必要检查，不能自动构成算法性能创新。分类头中CLIP段的权重平方范数占比也不能当作CLIP贡献了相同比例的判别信息；贡献需要通过配对预测和输入干预判断。")
    heading(document, "CLIP内容是否被实际利用", level=2)
    body_text(document, "我们将同一图像的CLIP特征替换为同域其他视频特征，或替换为统计量匹配的高斯噪声。固定正常模型的替换检查其输入依赖；采用同结构噪声重新训练，进一步检查收益是否仅来自新增参数和优化。替换时保留ViT输入，避免同时改变基线。")
    data_table(document, ["控制对照", "POOLED", "ADAPTIVE"], [
        ["其他视频替换 donor", "15／15", "9／15"],
        ["固定模型换噪声", "15／15", "10／15"],
        ["同结构噪声重训", "15／15", "0／15"]], [2.7, 1.4, 1.45])
    caption(document, "表3 真实输入优于自身控制且差值区间排除零的次数  3种子×5重复")
    body_text(document, "全局读取在三类控制中均显示真实图像CLIP内容的重要性，但这不等于它已经稳定超过强ViT基线。局部自适应读取没有超过噪声重训的明确证据，说明当前九区域聚合协议的可用增量尚未得到支持；不能据此推断所有局部CLIP信息都无用。表3重复共享目标数据，不是15个独立外部验证集。")
    heading(document, "修正幅度与迁移损伤", level=2)
    body_text(document, "复验还观察到，部分修正头在源域带来正向变化，却在跨域损害原有排序；增强修正惩罚或缩小修正幅度后，损伤减弱。自适应权重与均匀权重没有被证明存在稳定差异，因而目前不支持把失败主要归因于区域注意力过于集中。这个发现提示后续应检验修正的迁移稳定性和误伤，而不只增加融合层数。")
    body_text(document, "现阶段新增的同维度向量融合脚本提供固定权重、可学习全局权重、纯ViT、拼接及噪声／donor对照，已通过合成程序自检，尚未取得真实数据运行结果。它属于后续验证准备，不计为已实现的性能提升，也不代表已修复Bridge的多层信息利用。")
    heading(document, "六 前后端原型及使用边界")
    body_text(document, "当前原型包含前端界面与后端检测模块，后端使用改进单ViT，能够对输入图片进行真假鉴别并向界面返回结果。该功能使研究模型可以被实际调用，适合演示和辅助试验。原型与研究融合路线分开维护，避免在融合增益尚不确定时，把复杂研究模块描述为已经上线的有效功能。")
    body_text(document, "按当前使用情况，对更逼真的伪造图片，模型检测能力仍不及预期。未知生成算法、复杂人脸条件以及社媒压缩传播可能改变判别线索，系统输出需要人工复核，不具备自动作出身份真实性或司法证据结论的适用基础。本期没有完成真实业务数据的系统验收，也没有形成接口安全、并发吞吐、延迟或长期运行可靠性的量化结论。")
    body_text(document, "演示时应将图片鉴别流程与误判案例一并说明，不只展示正确样例。涉及人脸的图片应在授权范围内使用，不向未经授权的第三方传播；检测结果不能替代人工判断。后续若开展外部试用，应先明确输入范围、数据保护和错误处理规则，再评估推广。")
    heading(document, "七 项目成果与未完成目标")
    body_text(document, "本期形成的成果包括改进ViT模型实现与训练记录、已接入单ViT的图片鉴别前后端原型、研究用融合框架、受控诊断程序以及阶段分析资料。项目没有取得本项目论文发表、专利申请或授权成果；申报材料列举的既有导师专利和相关论文仅属于研究基础，不作为本项目新增成果计算。")
    body_text(document, "未完成目标主要包括：原计划空域—频域双向去偏的完整实现与增益确认，稳定超过独立改进ViT的跨域融合效果，重度压缩和高逼真样本的可靠检测，以及轻量化、毫秒级推理和高并发服务指标。没有完成的事项继续列为研究与工程任务，不用设计预期代替实测结果。当前项目的价值主要体现为可运行原型和对融合问题的具体定位，而不是已经达到部署级技术标准。")
    heading(document, "八 后续工作与结论")
    body_text(document, "下一步首先将部署模型、独立ViT原生分类头、重新训练的ViT-only头和研究融合方案放到同一输入与评测清单下比较，确认权重、预处理和标签一致。保存逐样本预测，分析CLIP修正救回了哪些错误、又误伤了哪些正确样本，再决定是否采用受约束修正或改进多层读取。若没有稳定互补证据，应保留简单单ViT原型，不以复杂结构替代有效验证。")
    body_text(document, "同时需要在未反复查看的新视频留出集上进行源域选模后的评估，并补齐高逼真伪造、降质输入和前后端运行测试。部分实验的逐样本数组与模型权重尚未同步至当前工作区，汇总可用于阶段展示，但进一步独立重算需取回原实验机存档。行政信息、模型版本、样本清单和演示证据也应在正式提交前逐项核对。")
    body_text(document, "总体而言，本项目已完成检测原型与多轮跨域融合验证，积累了关于当前实现缺陷和融合读取边界的证据。项目尚未实现预期的稳定跨域提升和全部工程成果，现阶段结论应限定在已验证的输入与实验协议内。后续工作将以真实增量和应用可靠性为依据推进，而不将阶段性能停滞解释为算法上限已被证明。")
    heading(document, "参考文献与实验取数说明")
    for reference_text in [
        "[1] Rössler A, Cozzolino D, Verdoliva L, et al. FaceForensics++ Learning to Detect Manipulated Facial Images. ICCV, 2019.",
        "[2] Li Y, Yang X, Sun P, et al. Celeb-DF A Large-scale Challenging Dataset for DeepFake Forensics. CVPR, 2020.",
        "[3] Radford A, Kim J W, Hallacy C, et al. Learning Transferable Visual Models From Natural Language Supervision. ICML, 2021.",
        "[4] Dosovitskiy A, Beyer L, Kolesnikov A, et al. An Image is Worth 16x16 Words Transformers for Image Recognition at Scale. ICLR, 2021.",
        "[5] acc.txt，第15至50轮训练记录。表1取第50轮，图2展示全部36轮。",
        "[6] vit_module/_g28/clip_readout_run01内的summary.json、config.json及comparisons.json，提供表2与表3数据。",
        "[7] vit_m2f2_detector_bridge.py、vit_adaptive_mattn_aps.py、train_bridge_phase1.py及项目工作记录。融合架构参考M2F2_Det工程。"]:
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, spacing=1.2, after=3)
        format_run(paragraph.add_run(reference_text), size=10.5)
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    memory = io.BytesIO()
    document.save(memory)
    editable = {"word/document.xml", "word/styles.xml", "word/settings.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    with ZipFile(REFERENCE) as source, ZipFile(memory) as generated, ZipFile(OUTPUT, "w") as destination:
        for info in source.infolist():
            destination.writestr(info, generated.read(info.filename) if info.filename in editable else source.read(info.filename))
        for info in generated.infolist():
            if info.filename not in source.namelist():
                destination.writestr(info, generated.read(info.filename))
    with ZipFile(REFERENCE) as source, ZipFile(OUTPUT) as result:
        preserved = [name for name in source.namelist() if name not in editable]
        assert all(source.read(name) == result.read(name) for name in preserved)
        all_text = Document(OUTPUT)._element.body.xpath(".//w:t/text()")
        combined = "\n".join(all_text)
        assert "98.162" in combined and "84.074" in combined and "未提出专利申请" in combined
        assert "B1-3请" not in combined and "负责人填首位" not in combined
        check = {"output": str(OUTPUT), "reference_sha256": expected_hash,
                 "output_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                 "preserved_parts": len(preserved), "new_parts": sorted(set(result.namelist()) - set(source.namelist())),
                 "sections": len(Document(OUTPUT).sections), "text_characters": len(combined),
                 "pending_fields": sorted(set(re.findall(r"【[^】]+】", combined))),
                 "baseline_macro": pooled["baseline_macro_auc"], "pooled_macro": pooled["macro_auc"]}
        (TASK_DIR / "build_checks.json").write_text(json.dumps(check, ensure_ascii=False, indent=2), encoding="utf-8")
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == expected_hash
    print(json.dumps(check, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
